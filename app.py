"""
5S Audit — Confection Textile
Backend Flask (API REST + rendu du front-end existant).

Lancer avec :  python app.py
Puis ouvrir :  http://127.0.0.1:5000
"""
import json
import base64
import binascii
from io import BytesIO
from datetime import datetime

from flask import Flask, jsonify, render_template, request
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from flask import send_file

from db import LINES, get_connection, init_db

app = Flask(__name__)


# ----------------------------------------------------------------------
# Page d'accueil (sert le front-end)
# ----------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html", lines=LINES)


# ----------------------------------------------------------------------
# Auth auditeur
# ----------------------------------------------------------------------
@app.route("/api/login", methods=["POST"])
def login():
    data = request.get_json(silent=True) or {}
    auditor_id = (data.get("id") or "").strip().upper()
    pin = (data.get("pin") or "").strip()

    conn = get_connection()
    row = conn.execute(
        "SELECT id, name, role FROM auditors WHERE id = ? AND pin = ?",
        (auditor_id, pin),
    ).fetchone()
    conn.close()

    if row is None:
        return jsonify({"ok": False, "error": "Identifiant ou code PIN incorrect."}), 401

    return jsonify({"ok": True, "auditor": {
        "id": row["id"], "name": row["name"], "role": row["role"] or "Auditeur 5S",
    }})


# ----------------------------------------------------------------------
# Écran ligne de production (Résultats)
# ----------------------------------------------------------------------
def _norm_line(section, line):
    """Cutting/finishing n'ont pas de ligne -> clé vide en base."""
    return line if section == "sewing" else ""


def _photo_list(value):
    """Lit les anciennes photos simples et les nouvelles listes JSON."""
    if not value:
        return []
    try:
        parsed = json.loads(value)
        return parsed if isinstance(parsed, list) else [value]
    except (TypeError, json.JSONDecodeError):
        return [value]


def _photo_json(value):
    photos = value if isinstance(value, list) else [value]
    return json.dumps([photo for photo in photos if photo], ensure_ascii=False)


def _text_list(value):
    """Lit les anciens commentaires simples et les nouveaux tableaux JSON."""
    if not value:
        return []
    try:
        parsed = json.loads(value)
        return parsed if isinstance(parsed, list) else [value]
    except (TypeError, json.JSONDecodeError):
        return [value]


def _text_json(value):
    comments = value if isinstance(value, list) else [value]
    return json.dumps([comment or "" for comment in comments], ensure_ascii=False)


def _has_text(value):
    values = value if isinstance(value, list) else [value]
    return any(str(comment or "").strip() for comment in values)


@app.route("/api/results")
def get_results():
    section = request.args.get("section", "sewing")
    line = _norm_line(section, request.args.get("line", "Line 01"))

    conn = get_connection()
    record = conn.execute(
        """SELECT * FROM audit_records
           WHERE section = ? AND line = ?
           ORDER BY id DESC LIMIT 1""",
        (section, line),
    ).fetchone()
    champion = conn.execute(
        "SELECT name, photo FROM champions WHERE section = ? AND line = ?",
        (section, line),
    ).fetchone()
    conn.close()

    label = f"Sewing — {line}" if section == "sewing" else section.capitalize()

    if record is None:
        return jsonify({
            "ok": True,
            "label": label,
            "record": None,
            "champion": {
                "name": champion["name"] if champion else "—",
                "photo": champion["photo"] if champion else None,
            },
        })

    return jsonify({
        "ok": True,
        "label": label,
        "record": {
            "id": record["id"],
            "points": record["points"],
            "rating": record["rating"],
            "before": _text_list(record["comment_before"]),
            "after": _text_list(record["comment_after"]),
            "improvement": record["improvement"],
            "photoBefore": _photo_list(record["photo_before"]),
            "photoAfter": _photo_list(record["photo_after"]),
        },
        "champion": {
            "name": champion["name"] if champion else "—",
            "photo": champion["photo"] if champion else None,
        },
    })


# ----------------------------------------------------------------------
# Nouvel audit
# ----------------------------------------------------------------------
@app.route("/api/audits", methods=["POST"])
def create_audit():
    data = request.get_json(silent=True) or {}

    section = data.get("section")
    if section not in ("cutting", "sewing", "finishing"):
        return jsonify({"ok": False, "error": "Section invalide."}), 400

    line = _norm_line(section, data.get("line", ""))
    if section == "sewing" and line not in LINES:
        return jsonify({"ok": False, "error": "Ligne invalide."}), 400

    photo_before = data.get("photoBefore")
    photo_after = data.get("photoAfter")
    comment_before = data.get("commentBefore")
    comment_after = data.get("commentAfter")
    if not _has_text(comment_before):
        return jsonify({"ok": False, "error": "Le commentaire Before est requis."}), 400
    if not _has_text(comment_after):
        return jsonify({"ok": False, "error": "Le commentaire After est requis."}), 400

    conn = get_connection()
    conn.execute(
        """INSERT INTO audit_records
           (section, line, auditor_id, photo_before, photo_after,
            comment_before, comment_after, rating, points, improvement)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            section,
            line,
            data.get("auditorId"),
            _photo_json(photo_before),
            _photo_json(photo_after),
            _text_json(data.get("commentBefore", "")),
            _text_json(data.get("commentAfter", "")),
            int(data.get("rating") or 0),
            int(data.get("points") or 0),
            data.get("improvement", ""),
        ),
    )
    conn.commit()
    conn.close()

    label = f"Sewing — {line}" if section == "sewing" else section.capitalize()
    return jsonify({"ok": True, "message": f"{label} : l'écran de la ligne a été mis à jour avec ce nouvel audit."})


@app.route("/api/audits")
def list_audits():
    conn = get_connection()
    rows = conn.execute(
        """SELECT a.*, au.name AS auditor_name, au.role AS auditor_role
           FROM audit_records a LEFT JOIN auditors au ON au.id = a.auditor_id
              WHERE a.auditor_id IS NOT NULL
           ORDER BY a.id DESC"""
    ).fetchall()
    conn.close()
    return jsonify({"ok": True, "audits": [
        {
            "id": row["id"], "section": row["section"], "line": row["line"],
            "auditorId": row["auditor_id"] or "—",
            "auditorName": row["auditor_name"] or row["auditor_id"] or "—",
            "auditorRole": row["auditor_role"] or "Auditeur 5S",
            "createdAt": row["created_at"],
            "photoBefore": _photo_list(row["photo_before"]),
            "photoAfter": _photo_list(row["photo_after"]),
            "commentBefore": _text_list(row["comment_before"]),
            "commentAfter": _text_list(row["comment_after"]),
            "rating": row["rating"], "points": row["points"],
            "improvement": row["improvement"] or "",
        }
        for row in rows
    ]})


@app.route("/api/audits/export")
def export_audits():
    conn = get_connection()
    rows = conn.execute(
        """SELECT a.*, au.name AS auditor_name, au.role AS auditor_role
           FROM audit_records a LEFT JOIN auditors au ON au.id = a.auditor_id
           WHERE a.auditor_id IS NOT NULL
           ORDER BY a.id DESC"""
    ).fetchall()
    conn.close()

    workbook = Workbook()
    workbook.remove(workbook.active)
    thin_gray = Side(style="thin", color="D7DEE8")
    card_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for row in rows:
        sheet = workbook.create_sheet(title=f"Audit {row['id']}")
        sheet.sheet_view.showGridLines = False
        for column in ("A", "B", "C", "D", "E", "F", "G", "H"):
            sheet.column_dimensions[column].width = 13
        for index in range(1, 14):
            sheet.row_dimensions[index].height = 22

        sheet.merge_cells("A1:H1")
        title = sheet["A1"]
        title.value = "Avant / Après"
        title.font = Font(name="Calibri", size=18, bold=False, color="26364A")
        title.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 32
        sheet["A1"].border = Border(bottom=Side(style="thin", color="D7DEE8"))

        sheet.merge_cells("A2:H2")
        metadata = sheet["A2"]
        metadata.value = (
            f"Publication #{row['id']}  |  {row['section'].capitalize()}"
            f" {row['line'] or ''}  |  {row['created_at']}  |  "
            f"{row['auditor_name'] or row['auditor_id'] or '-'} — "
            f"{row['auditor_role'] or 'Auditeur 5S'}"
        )
        metadata.font = Font(size=10, color="64748B")
        metadata.alignment = Alignment(vertical="center")

        sheet.merge_cells("A4:D4")
        sheet.merge_cells("E4:H4")
        before_label = sheet["A4"]
        before_label.value = "●  Before"
        before_label.font = Font(size=12, bold=True, color="C9232F")
        after_label = sheet["E4"]
        after_label.value = "●  After"
        after_label.font = Font(size=12, bold=True, color="55A936")

        before_photos = _photo_list(row["photo_before"])
        after_photos = _photo_list(row["photo_after"])
        _add_excel_photo(sheet, before_photos[0] if before_photos else None, "A5")
        _add_excel_photo(sheet, after_photos[0] if after_photos else None, "E5")
        for cell in ("A5", "E5"):
            sheet[cell].border = card_border
        sheet.merge_cells("A10:D10")
        sheet.merge_cells("E10:H10")
        sheet["A10"] = f"{len(before_photos)} photo{'s' if len(before_photos) != 1 else ''}"
        sheet["E10"] = f"{len(after_photos)} photo{'s' if len(after_photos) != 1 else ''}"
        for cell in ("A10", "E10"):
            sheet[cell].font = Font(size=10, color="64748B")
            sheet[cell].alignment = Alignment(horizontal="right")

        sheet.merge_cells("A11:D11")
        sheet.merge_cells("E11:H11")
        sheet.merge_cells("A12:D13")
        sheet.merge_cells("E12:H13")
        for cell, comments in (("A12", _text_list(row["comment_before"])), ("E12", _text_list(row["comment_after"]))):
            sheet[cell] = "COMMENTAIRE\n" + (comments[0] if comments else "Aucun commentaire")
            sheet[cell].font = Font(size=11, color="64748B")
            sheet[cell].alignment = Alignment(vertical="top", wrap_text=True)
            sheet[cell].fill = PatternFill("solid", fgColor="F8FAFC")
            sheet[cell].border = card_border
        sheet["A11"] = "COMMENTAIRE"
        sheet["E11"] = "COMMENTAIRE"
        for cell in ("A11", "E11"):
            sheet[cell].font = Font(size=9, color="B7791F", bold=True)
            sheet[cell].alignment = Alignment(vertical="bottom")

        sheet.merge_cells("A15:H15")
        sheet["A15"] = f"Note : {row['rating'] or 0} / 5    |    Points : {row['points'] or 0}    |    Point à améliorer : {row['improvement'] or '-'}"
        sheet["A15"].font = Font(size=10, color="475569")
        sheet["A15"].alignment = Alignment(wrap_text=True)

    if not rows:
        sheet = workbook.create_sheet(title="Aucun audit")
        sheet["A1"] = "Aucun audit publié"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"audits_5s_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _add_excel_photo(sheet, data_url, anchor):
    if not data_url or "," not in data_url:
        return
    try:
        image_data = base64.b64decode(data_url.split(",", 1)[1])
        image = ExcelImage(BytesIO(image_data))
        image.width = 385
        image.height = 260
        sheet.add_image(image, anchor)
        sheet.row_dimensions[5].height = 195
    except (ValueError, TypeError, binascii.Error):
        return


@app.route("/api/audits/<int:audit_id>", methods=["PUT", "DELETE"])
def manage_audit(audit_id):
    conn = get_connection()
    existing = conn.execute("SELECT id FROM audit_records WHERE id = ?", (audit_id,)).fetchone()
    if existing is None:
        conn.close()
        return jsonify({"ok": False, "error": "Audit introuvable."}), 404

    if request.method == "DELETE":
        conn.execute("DELETE FROM audit_records WHERE id = ?", (audit_id,))
        conn.commit()
        conn.close()
        return jsonify({"ok": True})

    data = request.get_json(silent=True) or {}
    fields = {
        "comment_before": _text_json(data.get("commentBefore", "")),
        "comment_after": _text_json(data.get("commentAfter", "")),
        "rating": int(data.get("rating") or 0),
        "points": int(data.get("points") or 0),
        "improvement": data.get("improvement", ""),
    }
    if data.get("photoBefore"):
        fields["photo_before"] = _photo_json(data["photoBefore"])
    if data.get("photoAfter"):
        fields["photo_after"] = _photo_json(data["photoAfter"])
    assignments = ", ".join(f"{key} = ?" for key in fields)
    conn.execute(f"UPDATE audit_records SET {assignments} WHERE id = ?", (*fields.values(), audit_id))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


# ----------------------------------------------------------------------
# Champions 5S
# ----------------------------------------------------------------------
@app.route("/api/champions")
def list_champions():
    conn = get_connection()
    rows = conn.execute("SELECT section, line, name, photo FROM champions").fetchall()
    conn.close()
    return jsonify({
        "ok": True,
        "champions": [
            {"section": r["section"], "line": r["line"], "name": r["name"], "photo": r["photo"]}
            for r in rows
        ],
    })


@app.route("/api/champions", methods=["POST"])
def save_champion():
    data = request.get_json(silent=True) or {}
    section = data.get("section")
    if section not in ("cutting", "sewing", "finishing"):
        return jsonify({"ok": False, "error": "Section invalide."}), 400

    line = _norm_line(section, data.get("line", ""))
    name = (data.get("name") or "").strip()
    photo = data.get("photo")  # peut être None -> on ne change pas la photo existante

    conn = get_connection()
    existing = conn.execute(
        "SELECT id, name, photo FROM champions WHERE section = ? AND line = ?",
        (section, line),
    ).fetchone()

    if existing is None:
        conn.execute(
            "INSERT INTO champions (section, line, name, photo) VALUES (?, ?, ?, ?)",
            (section, line, name or "Sans nom", photo),
        )
    else:
        new_name = name or existing["name"]
        new_photo = photo if photo else existing["photo"]
        conn.execute(
            "UPDATE champions SET name = ?, photo = ? WHERE id = ?",
            (new_name, new_photo, existing["id"]),
        )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


# ----------------------------------------------------------------------

    #app.run(debug=True)
if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)